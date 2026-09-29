from pathlib import Path
import io
from pypdf import PdfReader
from docx import Document

from app.utils.text import normalize_whitespace


class DocumentParser:
    @staticmethod
    def parse_bytes(filename: str, content: bytes) -> str:
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return DocumentParser._parse_pdf(content)
        if ext == ".docx":
            return DocumentParser._parse_docx(content)
        if ext in {".txt", ".md"}:
            return normalize_whitespace(content.decode("utf-8", errors="ignore"))
        raise ValueError("Unsupported file type. Use PDF, DOCX, TXT or MD.")

    @staticmethod
    def _parse_pdf(content: bytes) -> str:
        reader = PdfReader(io.BytesIO(content))
        pages = [(page.extract_text() or "") for page in reader.pages]
        return normalize_whitespace("\n".join(pages))

    @staticmethod
    def _parse_docx(content: bytes) -> str:
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return normalize_whitespace("\n".join(parts))
